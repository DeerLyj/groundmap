"""Build stable, local evidence coordinates for PDF and XLSX sources."""

from __future__ import annotations

import math
import tempfile
from datetime import date, datetime, time
from pathlib import Path


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def _display(value) -> str:
    value = _json_value(value)
    return "" if value is None else str(value).replace("\n", " ").replace("|", "\\|")


def build_workbook_evidence(source: Path) -> tuple[str, dict]:
    """Return a searchable formula appendix and a lossless formula/value view."""
    from openpyxl import load_workbook

    formulas_book = load_workbook(source, data_only=False, read_only=False)
    values_book = load_workbook(source, data_only=True, read_only=False)
    sheets = []
    items = []
    missing_cached_values = 0
    appendix = [
        "",
        "<!-- groundmap:workbook-evidence:start -->",
        "## 工作簿公式与定位",
        "",
    ]
    try:
        for formula_sheet, value_sheet in zip(
            formulas_book.worksheets, values_book.worksheets
        ):
            min_col, min_row, max_col, max_row = (
                formula_sheet.min_column,
                formula_sheet.min_row,
                formula_sheet.max_column,
                formula_sheet.max_row,
            )
            from openpyxl.utils import get_column_letter

            used_range = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{max_row}"
            )
            rows = []
            formula_cells = []
            appendix.extend([f"### {formula_sheet.title} ({used_range})", ""])
            header_row = max(
                range(min_row, min(max_row, min_row + 49) + 1),
                key=lambda row: sum(
                    isinstance(formula_sheet.cell(row, col).value, str)
                    and not str(formula_sheet.cell(row, col).value).startswith("=")
                    for col in range(min_col, max_col + 1)
                ),
            )
            headers = {
                col: _display(formula_sheet.cell(header_row, col).value)
                for col in range(min_col, max_col + 1)
            }

            for row_index in range(min_row, max_row + 1):
                values = []
                formulas = []
                row_text = []
                for col_index in range(min_col, max_col + 1):
                    formula_cell = formula_sheet.cell(row_index, col_index)
                    value_cell = value_sheet.cell(row_index, col_index)
                    formula = (
                        str(formula_cell.value)
                        if formula_cell.data_type == "f"
                        else None
                    )
                    value = _json_value(value_cell.value)
                    values.append(value)
                    formulas.append(formula)
                    coordinate = formula_cell.coordinate
                    if formula:
                        if value is None:
                            missing_cached_values += 1
                        formula_cells.append(
                            {
                                "cell": coordinate,
                                "formula": formula,
                                "value": value,
                                "locator": {
                                    "kind": "xlsx",
                                    "sheet": formula_sheet.title,
                                    "range": coordinate,
                                },
                            }
                        )
                    label = headers.get(col_index, "")
                    shown = f"{coordinate}"
                    if label and row_index != header_row:
                        shown += f" {label}"
                    shown += f"={_display(value)}"
                    if formula:
                        shown += f" [formula: {formula}]"
                    row_text.append(shown)

                row_range = (
                    f"{get_column_letter(min_col)}{row_index}:"
                    f"{get_column_letter(max_col)}{row_index}"
                )
                rows.append(
                    {"row": row_index, "values": values, "formulas": formulas}
                )
                if any(value is not None for value in values) or any(formulas):
                    items.append(
                        {
                            "text": " | ".join(row_text),
                            "locator": {
                                "kind": "xlsx",
                                "sheet": formula_sheet.title,
                                "range": row_range,
                            },
                        }
                    )

            if formula_cells:
                appendix.extend(
                    ["| 单元格 | 公式 | 计算值 |", "| --- | --- | --- |"]
                )
                appendix.extend(
                    f"| {cell['cell']} | `{_display(cell['formula'])}` | {_display(cell['value'])} |"
                    for cell in formula_cells
                )
                appendix.append("")
            else:
                appendix.extend(["（此工作表没有公式。）", ""])
            sheets.append(
                {
                    "name": formula_sheet.title,
                    "used_range": used_range,
                    "rows": rows,
                    "formula_cells": formula_cells,
                }
            )
    finally:
        formulas_book.close()
        values_book.close()

    issues = []
    if missing_cached_values:
        issues.append(
            f"{missing_cached_values} 个公式单元格没有缓存计算值；公式已保留，值需在 Excel 中重算"
        )
    appendix.append("<!-- groundmap:workbook-evidence:end -->")
    return "\n".join(appendix) + "\n", {
        "schema_version": 1,
        "modality": "workbook",
        "status": "degraded" if issues else "success",
        "sheets": sheets,
        "items": items,
        "issues": issues,
    }


def _pdf_word_block(word: dict) -> dict:
    return {
        "text": word["text"],
        "bbox": [word["x0"], word["top"], word["x1"], word["bottom"]],
        "confidence": None,
    }


def build_docling_evidence(document: dict) -> dict:
    """Convert Docling provenance into the same page/bbox evidence schema."""
    page_meta = document.get("pages", {})
    pages_by_number = {}
    for key, value in page_meta.items():
        page_number = int(value.get("page_no", key))
        size = value.get("size", {})
        pages_by_number[page_number] = {
            "page": page_number,
            "width": float(size.get("width", 0)),
            "height": float(size.get("height", 0)),
            "coordinate_space": "pdf_points_top_left",
            "extraction": "docling",
            "text": "",
            "blocks": [],
        }

    for item in [*document.get("texts", []), *document.get("tables", [])]:
        text = str(item.get("text") or item.get("orig") or "").strip()
        if not text and isinstance(item.get("data"), dict):
            text = " ".join(
                str(cell.get("text", "")).strip()
                for cell in item["data"].get("table_cells", [])
                if str(cell.get("text", "")).strip()
            )
        if not text:
            continue
        for provenance in item.get("prov", []):
            page_number = int(provenance["page_no"])
            page = pages_by_number.get(page_number)
            if page is None:
                continue
            bbox = provenance.get("bbox") or {}
            if bbox.get("coord_origin") == "BOTTOMLEFT":
                box = [
                    bbox.get("l"),
                    page["height"] - bbox.get("t", page["height"]),
                    bbox.get("r"),
                    page["height"] - bbox.get("b", 0),
                ]
            else:
                box = [bbox.get("l"), bbox.get("t"), bbox.get("r"), bbox.get("b")]
            page["blocks"].append(
                {"text": text, "bbox": box, "confidence": None}
            )

    pages = [pages_by_number[number] for number in sorted(pages_by_number)]
    issues = []
    for page in pages:
        page["text"] = " ".join(block["text"] for block in page["blocks"])
        if not page["blocks"]:
            issues.append(f"PDF 第 {page['page']} 页未提取到文字，需视觉复核")
    return {
        "schema_version": 1,
        "modality": "pdf",
        "status": "degraded" if issues else "success",
        "page_count": len(pages),
        "pages": pages,
        "items": [
            {
                "text": page["text"],
                "locator": {"kind": "pdf", "page": page["page"]},
                "blocks": page["blocks"],
            }
            for page in pages
            if page["text"]
        ],
        "issues": issues,
    }


def _ocr_block(block: dict, scale_x: float, scale_y: float) -> dict:
    points = block.get("bbox") or []
    xs = [point[0] for point in points]
    ys = [point[1] for point in points]
    bbox = (
        [min(xs) * scale_x, min(ys) * scale_y, max(xs) * scale_x, max(ys) * scale_y]
        if xs and ys
        else None
    )
    return {
        "text": block.get("text", ""),
        "bbox": bbox,
        "confidence": block.get("confidence"),
    }


def build_pdf_evidence(source: Path, ocr_images) -> dict:
    """Extract page text/bboxes; OCR image-only pages without uploading them."""
    import pdfplumber
    from PIL import Image

    pages = []
    pending_ocr = []
    issues = []
    with tempfile.TemporaryDirectory(prefix="groundmap-pdf-") as temp_dir:
        temp_root = Path(temp_dir)
        with pdfplumber.open(source) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                words = page.extract_words(use_text_flow=True, keep_blank_chars=False)
                if words:
                    blocks = [_pdf_word_block(word) for word in words]
                    pages.append(
                        {
                            "page": page_number,
                            "width": float(page.width),
                            "height": float(page.height),
                            "coordinate_space": "pdf_points_top_left",
                            "extraction": "pdf_text",
                            "text": " ".join(block["text"] for block in blocks),
                            "blocks": blocks,
                        }
                    )
                    continue

                image_path = temp_root / f"page-{page_number}.png"
                page.to_image(resolution=150).save(image_path, format="PNG")
                pages.append(
                    {
                        "page": page_number,
                        "width": float(page.width),
                        "height": float(page.height),
                        "coordinate_space": "pdf_points_top_left",
                        "extraction": "ocr",
                        "text": "",
                        "blocks": [],
                    }
                )
                pending_ocr.append((pages[-1], image_path))

        results = ocr_images([path for _, path in pending_ocr])
        for page, image_path in pending_ocr:
            ok, payload = results[image_path]
            if not ok:
                issues.append(f"PDF 第 {page['page']} 页 OCR 失败: {payload}")
                continue
            with Image.open(image_path) as image:
                scale_x = page["width"] / image.width
                scale_y = page["height"] / image.height
            blocks = [
                _ocr_block(block, scale_x, scale_y)
                for block in payload.get("blocks", [])
                if block.get("text")
            ]
            page["text"] = " ".join(block["text"] for block in blocks)
            page["blocks"] = blocks
            page["ocr"] = {
                "engine": payload.get("engine"),
                "engine_version": payload.get("engine_version"),
                "average_confidence": payload.get("average_confidence"),
                "needs_review": payload.get("needs_review", True),
            }
            if not blocks:
                issues.append(f"PDF 第 {page['page']} 页未识别到文字，需视觉复核")
            elif payload.get("needs_review"):
                issues.append(f"PDF 第 {page['page']} 页 OCR 置信度偏低，需视觉复核")

    return {
        "schema_version": 1,
        "modality": "pdf",
        "status": "degraded" if issues else "success",
        "page_count": len(pages),
        "pages": pages,
        "items": [
            {
                "text": page["text"],
                "locator": {"kind": "pdf", "page": page["page"]},
                "blocks": page["blocks"],
            }
            for page in pages
            if page["text"]
        ],
        "issues": issues,
    }
